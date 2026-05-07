import io
import os
import re
import zipfile
import tempfile
from pathlib import Path

import streamlit as st
from transformers import pipeline
import fitz  # PyMuPDF
import pandas as pd
from docx import Document

# ---------------------------------------------------
# CONFIG
# ---------------------------------------------------

MODEL_PATH = r"C:\models\privacy-filter"

SUPPORTED_TYPES = [
    "txt",
    "csv",
    "xlsx",
    "docx",
    "pdf",
]

# ---------------------------------------------------
# CUSTOM REGEX PATTERNS
# ---------------------------------------------------
CUSTOM_PATTERNS = {
    "IBAN": r"\bCH\d{2}\s?\d{4}\s?\d{4}\s?\d{4}\s?\d{4}\s?\d\b",
    "UID": r"\bCHE-\d{3}\.\d{3}\.\d{3}\b",
    "CREDIT_CARD": r"\b(?:\d[ -]*?){13,16}\b",
    "PHONE_SWISS": r"\b(?:\+41|0)\s?(?:\d{2}|\(\d{2}\))\s?\d{3}\s?\d{2}\s?\d{2}\b",
    "AHV_NUMBER": r"\b756\.\d{4}\.\d{4}\.\d{2}\b",
    "IP_ADDRESS": r"\b(?:\d{1,3}\.){3}\d{1,3}\b",
    "LICENSE_PLATE_CH": r"\b(?:ZH|BE|LU|AG|SG|ZG|BS|BL|SO|TI|VD|VS|GE|FR|GR|TG|NE|JU|OW|NW|UR|SZ|GL|SH|AR|AI)\s?\d{1,6}\b",
    "FINANCIAL_VALUE": r"\b(?:CHF|EUR|USD|Fr\.?)\s?-?\d+(?:[’'\.\s]\d{3})*(?:[,.]\d{2})?(?:\s?(?:%|Mio\.?|k))?\b",
    "PERCENTAGE": r"\b-?\d+(?:[,.]\d+)?\s?%\b",
    "ADDRESS": r"\b[A-ZÄÖÜ][A-Za-zÄÖÜäöüß\-]+(?:strasse|straße|weg|platz|gasse|allee|ring|street|avenue)\s+\d+[A-Za-z]?(?:,\s*\d{4}\s+[A-ZÄÖÜ][A-Za-zÄÖÜäöüß\-]+)?\b",
}

CONTEXT_VALUE_PATTERNS = {
    "CUSTOMER_NUMBER": r"\b(?:Kundennummer|Kunden-Nr\.?|Kunden Nr\.?|Kd\.-Nr\.?|KdNr|Kunden-ID|Kunde-ID)\s*[:#-]?\s*([A-Za-z0-9\-\/]+)",
    "EMPLOYEE_NUMBER": r"\b(?:Personalnummer|Personal-Nr\.?|MA-Nr\.?|Mitarbeiter-Nr\.?|Mitarbeiter ID|Mitarbeiter-ID)\s*[:#-]?\s*([A-Za-z0-9\-\/]+)",
    "CUSTOMER_REFERENCE": r"\b(?:Referenznummer|Referenz-Nr\.?|Ref\.?|Ticket-ID|Case-ID|Vorgangsnummer|Auftragsnummer|Bestellnummer|Order-ID)\s*[:#-]?\s*([A-Za-z0-9\-\/]+)",
    "INVOICE_NUMBER": r"\b(?:Rechnungsnummer|Rechnung-Nr\.?|Invoice-ID|Invoice Number|Belegnummer)\s*[:#-]?\s*([A-Za-z0-9\-\/]+)",
    "PAYMENT_REFERENCE": r"\b(?:QR-Referenz|ESR-Nummer|Referenznummer)\s*[:#-]?\s*([\d\s]+)",
    "BIRTHDATE": r"\b(?:Geburtsdatum|Geb\.?|DOB|Date of Birth)\s*[:#-]?\s*(\d{1,2}[.\-/]\d{1,2}[.\-/]\d{2,4})",
    "API_KEY": r"\b(?:api[_-]?key|token|secret)\s*[:=]\s*([A-Za-z0-9_\-]{16,})",
    "FINANCIAL_VALUE_CONTEXT": r"\b(?:Umsatz|Gewinn|Marge|Lohn|Gehalt|Salär|Bonus|Provision|EBITDA|Kosten|Budget)\s*[:#-]?\s*((?:CHF|EUR|USD|Fr\.?)\s?-?\d+(?:[’'\.\s]\d{3})*(?:[,.]\d{2})?(?:\s?(?:%|Mio\.?|k))?)",
}

# ---------------------------------------------------
# STREAMLIT CONFIG
# ---------------------------------------------------

st.set_page_config(
    page_title="Local Privacy Redactor",
    layout="wide",
)

# ---------------------------------------------------
# LOAD MODEL
# ---------------------------------------------------


@st.cache_resource
def load_classifier():

    return pipeline(
        "token-classification",
        model=MODEL_PATH,
        tokenizer=MODEL_PATH,
        aggregation_strategy="simple",
    )


classifier = load_classifier()

# ---------------------------------------------------
# DETECT ML ENTITIES
# ---------------------------------------------------


def detect_ml_entities(text, min_score=0.75):

    if not text.strip():
        return []

    results = classifier(text)

    cleaned = []

    for entity in results:
        word = entity.get("word", "").strip()

        if not word:
            continue

        if word in [".", ",", ";", ":"]:
            continue

        if entity.get("score", 0) < min_score:
            continue

        cleaned.append(entity)

    return cleaned


# ---------------------------------------------------
# DETECT REGEX ENTITIES
# ---------------------------------------------------


def detect_regex_entities(text):
    regex_entities = []

    for label, pattern in CUSTOM_PATTERNS.items():
        matches = re.finditer(pattern, text, flags=re.IGNORECASE)

        for match in matches:
            regex_entities.append({
                "start": match.start(),
                "end": match.end(),
                "word": match.group(),
                "entity_group": label,
                "score": 1.0,
            })

    for label, pattern in CONTEXT_VALUE_PATTERNS.items():
        matches = re.finditer(pattern, text, flags=re.IGNORECASE)

        for match in matches:
            value_start = match.start(1)
            value_end = match.end(1)
            value = match.group(1)

            regex_entities.append({
                "start": value_start,
                "end": value_end,
                "word": value,
                "entity_group": label,
                "score": 1.0,
            })

    return regex_entities


# ---------------------------------------------------
# MERGE ENTITIES
# ---------------------------------------------------


def merge_entities(text, entities):

    merged = []

    sorted_entities = sorted(entities, key=lambda x: x["start"])

    for entity in sorted_entities:
        if not merged:
            merged.append(entity.copy())
            continue

        last = merged[-1]

        overlap = entity["start"] <= last["end"]

        same_label = entity["entity_group"] == last["entity_group"]

        small_gap = entity["start"] - last["end"] <= 2

        if overlap or (same_label and small_gap):
            last["end"] = max(last["end"], entity["end"])

            last["word"] = text[last["start"] : last["end"]]

        else:
            merged.append(entity.copy())

    return merged


# ---------------------------------------------------
# REDACT TEXT
# ---------------------------------------------------


def redact_text(text, min_score=0.75):

    ml_entities = detect_ml_entities(text, min_score=min_score)

    regex_entities = detect_regex_entities(text)

    entities = ml_entities + regex_entities

    entities = merge_entities(text, entities)

    redacted_text = text

    for entity in sorted(entities, key=lambda x: x["start"], reverse=True):
        label = entity["entity_group"]

        replacement = f"[REDACTED_{label}]"

        redacted_text = (
            redacted_text[: entity["start"]]
            + replacement
            + redacted_text[entity["end"] :]
        )

    return redacted_text, entities


# ---------------------------------------------------
# TXT
# ---------------------------------------------------


def process_txt(uploaded_file, min_score):

    text = uploaded_file.read().decode("utf-8", errors="ignore")

    redacted, entities = redact_text(text, min_score)

    return (redacted.encode("utf-8"), entities, "text/plain", ".txt")


# ---------------------------------------------------
# CSV
# ---------------------------------------------------


def redact_cell(value, min_score, entities_all):

    if pd.isna(value):
        return value

    text = str(value)

    redacted, entities = redact_text(text, min_score)

    entities_all.extend(entities)

    return redacted


def process_csv(uploaded_file, min_score):

    df = pd.read_csv(uploaded_file)

    entities_all = []

    for col in df.columns:
        df[col] = df[col].apply(lambda x: redact_cell(x, min_score, entities_all))

    output = io.StringIO()

    df.to_csv(output, index=False)

    return (output.getvalue().encode("utf-8"), entities_all, "text/csv", ".csv")


# ---------------------------------------------------
# XLSX
# ---------------------------------------------------


def process_xlsx(uploaded_file, min_score):

    df = pd.read_excel(uploaded_file)

    entities_all = []

    for col in df.columns:
        df[col] = df[col].apply(lambda x: redact_cell(x, min_score, entities_all))

    output = io.BytesIO()

    df.to_excel(output, index=False)

    output.seek(0)

    return (
        output.read(),
        entities_all,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ".xlsx",
    )


# ---------------------------------------------------
# DOCX
# ---------------------------------------------------


def process_docx(uploaded_file, min_score):

    doc = Document(uploaded_file)

    entities_all = []

    for paragraph in doc.paragraphs:
        redacted, entities = redact_text(paragraph.text, min_score)

        entities_all.extend(entities)

        paragraph.text = redacted

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                redacted, entities = redact_text(cell.text, min_score)

                entities_all.extend(entities)

                cell.text = redacted

    output = io.BytesIO()

    doc.save(output)

    output.seek(0)

    return (
        output.read(),
        entities_all,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".docx",
    )


# ---------------------------------------------------
# PDF
# ---------------------------------------------------


def process_pdf(uploaded_file, min_score):

    entities_all = []

    with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
        tmp.write(uploaded_file.read())

        tmp_path = tmp.name

    doc = fitz.open(tmp_path)

    for page in doc:
        text = page.get_text()

        _, entities = redact_text(text, min_score)

        entities_all.extend(entities)

        for entity in entities:
            sensitive_text = text[entity["start"] : entity["end"]].strip()

            if not sensitive_text:
                continue

            rectangles = page.search_for(sensitive_text)

            for rect in rectangles:
                page.add_redact_annot(rect, fill=(0, 0, 0))

        page.apply_redactions()

    output = io.BytesIO()

    doc.save(output)

    doc.close()

    os.remove(tmp_path)

    output.seek(0)

    return (output.read(), entities_all, "application/pdf", ".pdf")


# ---------------------------------------------------
# FILE ROUTER
# ---------------------------------------------------


def process_file(uploaded_file, min_score):

    suffix = Path(uploaded_file.name).suffix.lower()

    if suffix == ".txt":
        return process_txt(uploaded_file, min_score)

    if suffix == ".csv":
        return process_csv(uploaded_file, min_score)

    if suffix == ".xlsx":
        return process_xlsx(uploaded_file, min_score)

    if suffix == ".docx":
        return process_docx(uploaded_file, min_score)

    if suffix == ".pdf":
        return process_pdf(uploaded_file, min_score)

    raise ValueError(f"Nicht unterstütztes Format: {suffix}")


# ---------------------------------------------------
# UI
# ---------------------------------------------------

st.title("🔒 Local Privacy Redactor")

st.write("Dokumente lokal hochladen, redacten und sicher herunterladen.")

st.caption("100% lokal • Keine Cloud • Keine API")

with st.sidebar:
    st.header("Einstellungen")

    min_score = st.slider("Confidence Threshold", 0.0, 1.0, 0.75, 0.05)

uploaded_files = st.file_uploader(
    "Dokumente hochladen",
    type=SUPPORTED_TYPES,
    accept_multiple_files=True,
)

# ---------------------------------------------------
# PROCESS FILES
# ---------------------------------------------------

if uploaded_files:
    redacted_files = []

    for uploaded_file in uploaded_files:
        st.divider()

        st.subheader(uploaded_file.name)

        try:
            with st.spinner("Lokale Redaction läuft..."):
                data, entities, mime, suffix = process_file(uploaded_file, min_score)

            output_name = Path(uploaded_file.name).stem + "_redacted" + suffix

            redacted_files.append((output_name, data))

            col1, col2 = st.columns(2)

            with col1:
                st.metric("Gefundene Entities", len(entities))

            with col2:
                st.download_button(
                    label="⬇ Download",
                    data=data,
                    file_name=output_name,
                    mime=mime,
                    key=uploaded_file.name,
                )

            if entities:
                with st.expander("Gefundene sensitive Daten"):
                    preview = pd.DataFrame([
                        {
                            "Text": e["word"],
                            "Typ": e["entity_group"],
                            "Score": round(float(e["score"]), 3),
                        }
                        for e in entities
                    ])

                    st.dataframe(preview, use_container_width=True)

        except Exception as e:
            st.error(f"Fehler bei {uploaded_file.name}: {e}")

    # ---------------------------------------------------
    # ZIP DOWNLOAD
    # ---------------------------------------------------

    if len(redacted_files) > 1:
        zip_buffer = io.BytesIO()

        with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
            for filename, file_data in redacted_files:
                zip_file.writestr(filename, file_data)

        zip_buffer.seek(0)

        st.download_button(
            label="⬇ Alle Dateien als ZIP herunterladen",
            data=zip_buffer,
            file_name="redacted_documents.zip",
            mime="application/zip",
        )
